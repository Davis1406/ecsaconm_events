import io
import secrets
from typing import Annotated, Optional
from fastapi import APIRouter, BackgroundTasks, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel
from sqlalchemy.orm import Session, joinedload

from core.database import get_db
from dependencies.auth_dependency import Auth, get_current_user
from models.models import (
    DepartureDetail,
    Event,
    ParticipationRole,
    Registration,
    SystemSetting,
    User,
)

router = APIRouter()

user_dependency = Annotated[dict, Depends(get_current_user)]

DIGEST_BATCH_SIZE = 10
VIEW_TOKEN_SETTING_KEY = "departure_details_view_token"
DEFAULT_EVENT_ID = 1


def get_auth_dep(db: Session = Depends(get_db)) -> Auth:
    return Auth(db)


def _get_or_create_view_token(db: Session) -> str:
    """A single standing secret that lets lemmym@/info@ open the live
    submissions list by URL alone, no login — generated once and reused
    (not per-person, since this link is shared by two people)."""
    row = db.query(SystemSetting).filter_by(key=VIEW_TOKEN_SETTING_KEY).first()
    if row and row.value:
        return row.value
    token = secrets.token_urlsafe(32)
    if row:
        row.value = token
    else:
        row = SystemSetting(key=VIEW_TOKEN_SETTING_KEY, value=token)
        db.add(row)
    db.commit()
    return token


def _eligible_registration(db: Session, event_id: int, email: str):
    """The paid, non-secretariat registration matching this email for this
    event, or None. Secretariat is excluded outright regardless of its
    always-true is_paid status; everyone else must actually be paid."""
    email = (email or "").strip().lower()
    if not email:
        return None
    return (
        db.query(Registration)
        .join(Registration.user)
        .filter(
            Registration.deleted_at == None,
            Registration.event_id == event_id,
            Registration.paid == True,
            Registration.participation_role != ParticipationRole.secretariat,
            User.email.ilike(email),
        )
        .first()
    )


def _eligible_registrants(db: Session, event_id: Optional[int] = None):
    """Every paid, non-secretariat registrant — the audience this form's
    invitation email is sent to."""
    q = (
        db.query(Registration)
        .join(Registration.user)
        .options(joinedload(Registration.user))
        .filter(
            Registration.deleted_at == None,
            Registration.paid == True,
            Registration.participation_role != ParticipationRole.secretariat,
        )
    )
    if event_id:
        q = q.filter(Registration.event_id == event_id)

    results = []
    seen_emails = set()
    for r in q.all():
        user = r.user
        if not user or not user.email:
            continue
        email = user.email.strip().lower()
        if email in seen_emails:
            continue
        seen_emails.add(email)
        results.append({
            "email": email,
            "name": f"{user.firstname or ''} {user.lastname or ''}".strip() or email,
            "registration_id": r.id,
            "event_id": r.event_id,
        })
    return results


def _serialize(rec: DepartureDetail):
    return {
        "id": rec.id,
        "event_id": rec.event_id,
        "email": rec.email,
        "name": rec.name,
        "hotel": rec.hotel,
        "departure_date": rec.departure_date,
        "departure_time": rec.departure_time,
        "submitted": rec.submitted_at is not None,
        "submitted_at": rec.submitted_at.isoformat() if rec.submitted_at else None,
        "created_at": rec.created_at.isoformat() if rec.created_at else None,
    }


def _build_invitation_email(event_name, form_link):
    subject = f"Please Share Your Travel & Hotel Details — {event_name}"
    html = (
        f"<p>Dear Delegate,</p>"
        f"<p>As we finalise logistics for <strong>{event_name}</strong>, kindly let us know "
        f"your hotel and departure details using the short form below.</p>"
        f"<p><a href=\"{form_link}\" style=\"display:inline-block;padding:12px 28px;"
        f"background-color:rgb(254,80,103);color:#ffffff;text-decoration:none;"
        f"border-radius:8px;font-weight:600;\">Submit Travel Details</a></p>"
        f"<p>If the button doesn't work, copy and paste this link into your browser:<br>"
        f"<span style=\"color:#6b7280;\">{form_link}</span></p>"
        f"<p>Thank you,<br>ECSACONM Secretariat</p>"
    )
    return subject, html


def _build_receipt_email(event_name, name, hotel, departure_date, departure_time):
    subject = f"Received — Your Travel Details for {event_name}"
    html = (
        f"<p>Dear {name},</p>"
        f"<p>Thank you — we've received your travel details:</p>"
        f"<table style=\"border-collapse:collapse;font-size:14px;\">"
        f"<tr><td style=\"padding:4px 12px 4px 0;color:#6b7280;\">Hotel</td><td><strong>{hotel or '—'}</strong></td></tr>"
        f"<tr><td style=\"padding:4px 12px 4px 0;color:#6b7280;\">Departure date</td><td><strong>{departure_date or '—'}</strong></td></tr>"
        f"<tr><td style=\"padding:4px 12px 4px 0;color:#6b7280;\">Departure time</td><td><strong>{departure_time or '—'}</strong></td></tr>"
        f"</table>"
        f"<p>If any of this changes, just submit the form again with the same email address.</p>"
        f"<p>Thank you,<br>ECSACONM Secretariat</p>"
    )
    return subject, html


def _build_digest_email(event_name, entries, total, view_link):
    subject = f"Travel Details — {len(entries)} New Submissions ({event_name})"
    rows = "".join(
        f"<tr><td style='padding:4px 10px;border-bottom:1px solid #eee;'>{e.name or e.email}</td>"
        f"<td style='padding:4px 10px;border-bottom:1px solid #eee;'>{e.hotel or '—'}</td>"
        f"<td style='padding:4px 10px;border-bottom:1px solid #eee;'>{e.departure_date or '—'}</td>"
        f"<td style='padding:4px 10px;border-bottom:1px solid #eee;'>{e.departure_time or '—'}</td></tr>"
        for e in entries
    )
    html = (
        f"<p>{len(entries)} more registrants have submitted their travel details for "
        f"<strong>{event_name}</strong> ({total} total submitted so far).</p>"
        f"<table style=\"border-collapse:collapse;font-size:13px;width:100%;\">"
        f"<tr style=\"text-align:left;color:#6b7280;\"><th style='padding:4px 10px;'>Name</th>"
        f"<th style='padding:4px 10px;'>Hotel</th><th style='padding:4px 10px;'>Departure date</th>"
        f"<th style='padding:4px 10px;'>Departure time</th></tr>{rows}</table>"
        f"<p style=\"margin-top:16px;\">View the full, live list any time — no login needed:<br>"
        f"<a href=\"{view_link}\">{view_link}</a></p>"
    )
    return subject, html


@router.get("/recipients")
def list_recipients(
    current_user: user_dependency,
    db: Session = Depends(get_db),
    auth_dependency: Auth = Depends(get_auth_dep),
    event_id: int = Query(None),
):
    """Admin view: eligible (paid, non-secretariat) registrants and whether
    they've submitted yet."""
    auth_dependency.secure_access("ADMIN_DASHBOARD", current_user["user_id"])
    people = _eligible_registrants(db, event_id)
    out = []
    for p in people:
        rec = (
            db.query(DepartureDetail)
            .filter(
                DepartureDetail.email == p["email"],
                DepartureDetail.event_id == p["event_id"],
                DepartureDetail.deleted_at == None,
            )
            .first()
        )
        out.append({**p, "submitted": bool(rec and rec.submitted_at)})
    return out


@router.get("/list")
def list_submissions(
    current_user: user_dependency,
    db: Session = Depends(get_db),
    auth_dependency: Auth = Depends(get_auth_dep),
    event_id: int = Query(None),
):
    """Admin view: every submitted travel-details record, for the
    Registrations page's own Travel Details panel (list + export) — same
    data as the public-view link, but behind a login instead of a token."""
    auth_dependency.secure_access("ADMIN_DASHBOARD", current_user["user_id"])
    q = db.query(DepartureDetail).filter(DepartureDetail.deleted_at == None)
    if event_id:
        q = q.filter(DepartureDetail.event_id == event_id)
    records = q.order_by(DepartureDetail.submitted_at.desc().nullslast(), DepartureDetail.created_at.desc()).all()
    return {
        "total_submitted": sum(1 for r in records if r.submitted_at),
        "data": [_serialize(r) for r in records],
    }


@router.delete("/{detail_id}")
def delete_submission(
    detail_id: int,
    current_user: user_dependency,
    db: Session = Depends(get_db),
    auth_dependency: Auth = Depends(get_auth_dep),
):
    """Admin: remove one submission (e.g. a test entry) — soft delete, same
    convention as the rest of the app. The registrant can simply submit the
    form again afterwards if this was a real, mistakenly-removed entry."""
    auth_dependency.secure_access("ADMIN_DASHBOARD", current_user["user_id"])
    rec = db.query(DepartureDetail).filter(
        DepartureDetail.id == detail_id, DepartureDetail.deleted_at == None,
    ).first()
    if not rec:
        raise HTTPException(status_code=404, detail="Submission not found.")
    from datetime import datetime, timezone
    rec.deleted_at = datetime.now(timezone.utc)
    db.commit()
    return {"detail": "Submission deleted."}


@router.get("/export")
def export_submissions(
    current_user: user_dependency,
    db: Session = Depends(get_db),
    auth_dependency: Auth = Depends(get_auth_dep),
    event_id: int = Query(None),
):
    """Admin: export every submitted travel-details record to Excel, same
    layout convention as the Registrations export."""
    auth_dependency.secure_access("EXPORT_REGISTRATIONS", current_user["user_id"])
    q = db.query(DepartureDetail).filter(
        DepartureDetail.deleted_at == None, DepartureDetail.submitted_at != None,
    )
    if event_id:
        q = q.filter(DepartureDetail.event_id == event_id)
    records = q.order_by(DepartureDetail.submitted_at.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Travel Details"

    header_fill = PatternFill("solid", start_color="0095B6")
    alt_fill = PatternFill("solid", start_color="E8F4F8")
    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    body_font = Font(name="Arial", size=10)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    headers = ["#", "Name", "Email", "Hotel", "Departure Date", "Departure Time", "Submitted At"]
    ws.row_dimensions[1].height = 22
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(1, ci, h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
    ws.freeze_panes = "A2"

    for ri, r in enumerate(records, 2):
        use_fill = alt_fill if ri % 2 == 0 else PatternFill("solid", start_color="FFFFFF")
        row = [
            r.id, r.name, r.email, r.hotel, r.departure_date, r.departure_time,
            r.submitted_at.strftime("%d %b %Y %H:%M") if r.submitted_at else "",
        ]
        for ci, val in enumerate(row, 1):
            cell = ws.cell(ri, ci, val)
            cell.font = body_font
            cell.fill = use_fill
            cell.alignment = left

    col_widths = [6, 22, 30, 26, 16, 16, 18]
    for ci, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=travel_details_export.xlsx"},
    )


class SendFormBody(BaseModel):
    event_id: Optional[int] = None
    selected_emails: Optional[list] = None
    test_email: Optional[str] = None


@router.post("/send")
def send_departure_invitations(
    current_user: user_dependency,
    db: Session = Depends(get_db),
    auth_dependency: Auth = Depends(get_auth_dep),
    background_tasks: BackgroundTasks = BackgroundTasks(),
    body: SendFormBody = None,
):
    """Email the one shared public form link to each selected (or all)
    eligible registrant. Pass `test_email` to fire a single trial send
    immediately (not queued/logged as a bulk job)."""
    auth_dependency.secure_access("ADMIN_DASHBOARD", current_user["user_id"])

    body = body or SendFormBody()
    import utils.mailer_util as mailer_util

    event_id = body.event_id or DEFAULT_EVENT_ID
    event = db.query(Event).filter(Event.id == event_id).first()
    event_name = event.event if event else "ECSACONM Scientific Conference"
    form_link = f"{mailer_util.CLIENT_ORIGIN}/#/travel-details?event_id={event_id}"

    if body.test_email:
        subject, html = _build_invitation_email(event_name, form_link)
        mailer_util.send_email(body.test_email, subject, html, email_type="departure_details_invitation", sent_by_user_id=current_user["user_id"])
        return {"sent": 1, "message": f"Sample invitation sent to {body.test_email}."}

    selected_set = (
        {e.strip().lower() for e in body.selected_emails} if body.selected_emails else None
    )
    people = _eligible_registrants(db, event_id)
    subject, html = _build_invitation_email(event_name, form_link)

    jobs = []
    for p in people:
        if selected_set is not None and p["email"] not in selected_set:
            continue
        jobs.append({
            "recipient_email": p["email"],
            "subject": subject,
            "email_body": html,
            "email_type": "departure_details_invitation",
            "sent_by_user_id": current_user["user_id"],
        })

    sent = len(jobs)
    if jobs:
        background_tasks.add_task(mailer_util.send_bulk_emails, jobs)

    return {
        "sent": sent,
        "form_link": form_link,
        "message": f"Travel-details form link queued for {sent} paid, non-secretariat registrant(s).",
    }


@router.get("/form-link")
def get_form_link(
    current_user: user_dependency,
    db: Session = Depends(get_db),
    auth_dependency: Auth = Depends(get_auth_dep),
    event_id: int = Query(None),
):
    """Admin: the one shared public URL for the travel-details form."""
    auth_dependency.secure_access("ADMIN_DASHBOARD", current_user["user_id"])
    import utils.mailer_util as mailer_util
    event_id = event_id or DEFAULT_EVENT_ID
    return {"link": f"{mailer_util.CLIENT_ORIGIN}/#/travel-details?event_id={event_id}"}


class SubmitFormBody(BaseModel):
    name: str
    email: str
    hotel: str
    departure_date: str
    departure_time: str
    event_id: Optional[int] = None


@router.post("/submit")
def submit_departure_form(
    body: SubmitFormBody,
    db: Session = Depends(get_db),
    background_tasks: BackgroundTasks = BackgroundTasks(),
):
    """Public, no-auth: the one shared form's submit endpoint. Only accepts
    submissions from paid, non-secretariat registrants — matched by email —
    which is how "excluding secretariat" is enforced on a link anyone can
    open. Emails the submitter a receipt, and — every DIGEST_BATCH_SIZE
    submissions — sends lemmym@/info@ a batch digest instead of pinging
    them on every single one."""
    event_id = body.event_id or DEFAULT_EVENT_ID
    email = (body.email or "").strip().lower()

    reg = _eligible_registration(db, event_id, email)
    if not reg:
        raise HTTPException(
            status_code=403,
            detail="We couldn't match that email to a paid registration for this event. "
                   "Please use the email address you registered with, or contact the secretariat.",
        )

    from datetime import datetime, timezone
    rec = (
        db.query(DepartureDetail)
        .filter(
            DepartureDetail.event_id == event_id,
            DepartureDetail.email == email,
            DepartureDetail.deleted_at == None,
        )
        .first()
    )
    if not rec:
        rec = DepartureDetail(
            event_id=event_id,
            registration_id=reg.id,
            email=email,
            token=secrets.token_urlsafe(24),
        )
        db.add(rec)

    rec.name = (body.name or "").strip() or rec.name
    rec.hotel = (body.hotel or "").strip()
    rec.departure_date = (body.departure_date or "").strip()
    rec.departure_time = (body.departure_time or "").strip()
    rec.submitted_at = datetime.now(timezone.utc)
    db.commit()

    import utils.mailer_util as mailer_util
    event = db.query(Event).filter(Event.id == event_id).first()
    event_name = event.event if event else "ECSACONM Event"

    subject, html = _build_receipt_email(event_name, rec.name or rec.email, rec.hotel, rec.departure_date, rec.departure_time)
    background_tasks.add_task(
        mailer_util.send_email, rec.email, subject, html,
        "departure_details_receipt", None,
    )

    total_submitted = (
        db.query(DepartureDetail)
        .filter(
            DepartureDetail.event_id == event_id,
            DepartureDetail.deleted_at == None,
            DepartureDetail.submitted_at != None,
        )
        .count()
    )
    if total_submitted % DIGEST_BATCH_SIZE == 0:
        batch = (
            db.query(DepartureDetail)
            .filter(
                DepartureDetail.event_id == event_id,
                DepartureDetail.deleted_at == None,
                DepartureDetail.submitted_at != None,
            )
            .order_by(DepartureDetail.submitted_at.desc())
            .limit(DIGEST_BATCH_SIZE)
            .all()
        )
        token_v = _get_or_create_view_token(db)
        view_link = f"{mailer_util.CLIENT_ORIGIN}/#/travel-details-report/{token_v}?event_id={event_id}"
        subject, html = _build_digest_email(event_name, list(reversed(batch)), total_submitted, view_link)
        for to in ("lemmym@ecsaconm.org", "info@ecsaconm.org"):
            background_tasks.add_task(
                mailer_util.send_email, to, subject, html,
                "departure_details_digest", None,
            )

    return {
        "message": "Thank you! Your travel details have been recorded.",
        "hotel": rec.hotel,
        "departure_date": rec.departure_date,
        "departure_time": rec.departure_time,
    }


@router.get("/public-view/{token}")
def public_view(
    token: str,
    db: Session = Depends(get_db),
    event_id: int = Query(None),
):
    """Public, no-auth: the live submissions list, reachable only by whoever
    has this unguessable link (sent to lemmym@/info@) — same trust model as
    the app's other public-by-link views (room-view, document QR, etc.)."""
    real_token = _get_or_create_view_token(db)
    if token != real_token:
        raise HTTPException(status_code=404, detail="Not found.")

    q = db.query(DepartureDetail).filter(DepartureDetail.deleted_at == None)
    if event_id:
        q = q.filter(DepartureDetail.event_id == event_id)
    records = q.order_by(DepartureDetail.submitted_at.desc().nullslast(), DepartureDetail.created_at.desc()).all()

    event = db.query(Event).filter(Event.id == event_id).first() if event_id else None
    return {
        "event_name": event.event if event else None,
        "total_recipients": len(records),
        "total_submitted": sum(1 for r in records if r.submitted_at),
        "data": [_serialize(r) for r in records],
    }


@router.get("/view-link")
def get_view_link(
    current_user: user_dependency,
    db: Session = Depends(get_db),
    auth_dependency: Auth = Depends(get_auth_dep),
    event_id: int = Query(None),
):
    """Admin: the standing no-login URL to hand to lemmym@/info@."""
    auth_dependency.secure_access("ADMIN_DASHBOARD", current_user["user_id"])
    import utils.mailer_util as mailer_util
    token = _get_or_create_view_token(db)
    link = f"{mailer_util.CLIENT_ORIGIN}/#/travel-details-report/{token}"
    if event_id:
        link += f"?event_id={event_id}"
    return {"link": link}


@router.post("/sample-receipt")
def send_sample_receipt(
    current_user: user_dependency,
    db: Session = Depends(get_db),
    auth_dependency: Auth = Depends(get_auth_dep),
    to_email: str = Query(...),
    event_id: int = Query(None),
):
    """Admin: fire a one-off sample of the post-submission receipt email
    using placeholder data, so the wording/format can be reviewed."""
    auth_dependency.secure_access("ADMIN_DASHBOARD", current_user["user_id"])
    import utils.mailer_util as mailer_util
    event = db.query(Event).filter(Event.id == event_id).first() if event_id else db.query(Event).first()
    event_name = event.event if event else "ECSACONM Scientific Conference"
    subject, html = _build_receipt_email(event_name, "Jane Sample Delegate", "Sample Grand Hotel", "2026-09-20", "10:30")
    mailer_util.send_email(to_email, subject, html, email_type="departure_details_receipt", sent_by_user_id=current_user["user_id"])
    return {"sent": 1, "message": f"Sample receipt sent to {to_email}."}
